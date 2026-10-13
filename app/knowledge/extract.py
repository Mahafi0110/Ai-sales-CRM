"""Reads the text out of an uploaded file. Supports .txt, .md, .pdf and .docx. Files are not kept: only the
extracted text is stored."""
import io
import zipfile
from pathlib import PurePath

ALLOWED_EXTENSIONS = (".txt", ".md", ".markdown", ".pdf", ".docx")
MAX_PDF_PAGES = 300


class ExtractError(ValueError):
    """The file could not be read; the message is safe to show to the person."""


def _text_file(data: bytes) -> str:
    if b"\x00" in data[:2000]:
        raise ExtractError("This looks like a binary file, not text.")
    return data.decode("utf-8-sig", errors="replace")


def _pdf(data: bytes) -> str:
    from pypdf import PdfReader
    from pypdf.errors import PyPdfError
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted and not reader.decrypt(""):
            raise ExtractError("This PDF is password protected. Remove the password and upload it again.")
        pages = reader.pages[:MAX_PDF_PAGES]
        return "\n\n".join((page.extract_text() or "") for page in pages)
    except ExtractError:
        raise
    except (PyPdfError, ValueError, KeyError, OSError, RecursionError) as exc:
        raise ExtractError("This PDF could not be read. Try exporting it again, or paste the text instead.") from exc


def _docx(data: bytes) -> str:
    from docx import Document
    try:
        document = Document(io.BytesIO(data))
    except (zipfile.BadZipFile, KeyError, ValueError, OSError) as exc:
        raise ExtractError("This Word file could not be read. Is it a real .docx file?") from exc
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n\n".join(parts)


def extract_text(filename: str, data: bytes) -> str:
    extension = PurePath(filename or "").suffix.lower()
    if extension not in ALLOWED_EXTENSIONS:
        raise ExtractError("Please upload a .txt, .md, .pdf or .docx file, or paste the text.")
    text = {".pdf": _pdf, ".docx": _docx}.get(extension, _text_file)(data)
    if not text.strip():
        raise ExtractError("No text was found in this file (a scanned PDF contains only pictures). "
                           "Paste the text instead.")
    return text
